import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
df=pd.read_pickle('data/build/clean.pkl')
wb=Workbook()
F=lambda **k: Font(name='Arial',**k)
hdr_fill=PatternFill('solid',fgColor='7A003C'); flag_fill=PatternFill('solid',fgColor='FFF2CC')
thin=Border(bottom=Side(style='thin',color='DDD9E2'))

ws=wb.active; ws.title='Menu'
cols=list(df.columns)
ws.append(cols)
for c in ws[1]:
    c.font=F(bold=True,color='FFFFFF'); c.fill=hdr_fill; c.alignment=Alignment(vertical='center',wrap_text=True)
ws.row_dimensions[1].height=32
numcols=cols[5:18]
for _,r in df.iterrows():
    row=[]
    for c in cols:
        v=r[c]
        row.append(None if pd.isna(v) or v=='' else (float(v) if c in numcols else str(v)))
    ws.append(row)
n=len(df)+1
fmt={'Price':'$0.00','Calories':'#,##0','Sodium (mg)':'#,##0','Cholesterol (mg)':'#,##0','Calcium (mg)':'#,##0'}
for j,c in enumerate(cols,1):
    L=get_column_letter(j)
    for i in range(2,n+1):
        cell=ws[f'{L}{i}']; cell.font=F()
        if c in numcols: cell.number_format=fmt.get(c,'0.0#')
    w={'Location':22,'Station':26,'Category':20,'Menu Item':40,'Serving Size':12,'Data Note':44}.get(c,11)
    ws.column_dimensions[L].width=w
note_col=get_column_letter(cols.index('Data Note')+1)
for i in range(2,n+1):
    if ws[f'{note_col}{i}'].value:
        for c in ws[i]: c.fill=flag_fill
ws.freeze_panes='E2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(cols))}{n}'

# Other restaurants (non-McMaster franchises)
import json
FR=json.load(open('data/build/franchise.json'))
o=wb.create_sheet('Other Restaurants')
ocols=['Location','Restaurant','Category','Menu Item','Serving','Calories','Fat (g)','Saturated Fat (g)','Cholesterol (mg)','Sodium (mg)','Carbohydrate (g)','Fibre (g)','Sugars (g)','Protein (g)','Type','Source']
o.append(ocols)
for c in o[1]: c.font=F(bold=True,color='FFFFFF'); c.fill=hdr_fill; c.alignment=Alignment(vertical='center',wrap_text=True)
TYPE={0:'Meal',1:'Add-on',2:'Snack or drink'}
for r in FR:
    o.append([r['location'],r['station'],r['category'],r['item'],r['serving'],r['calories'],r['fat'],r['satfat'],r['chol'],r['sodium'],r['carbs'],r['fibre'],r['sugars'],r['protein'],TYPE[r['kind']],r['source']])
for j,c in enumerate(ocols,1):
    L=get_column_letter(j)
    o.column_dimensions[L].width={'Location':14,'Restaurant':14,'Category':22,'Menu Item':46,'Serving':11,'Type':14,'Source':50}.get(c,11)
    for i in range(2,len(FR)+2):
        o[f'{L}{i}'].font=F()
        if 6<=j<=14: o[f'{L}{i}'].number_format='#,##0' if c in ('Calories','Sodium (mg)','Cholesterol (mg)') else '0.0#'
o.freeze_panes='E2'; o.auto_filter.ref=f'A1:{get_column_letter(len(ocols))}{len(FR)+1}'

# Summary with formulas
s=wb.create_sheet('By Location')
s.append(['Location','Menu items','Avg calories','Avg protein (g)','Avg carbs (g)','Avg fat (g)'])
for c in s[1]: c.font=F(bold=True,color='FFFFFF'); c.fill=hdr_fill
col=lambda name: f"Menu!${get_column_letter(cols.index(name)+1)}$2:${get_column_letter(cols.index(name)+1)}${n}"
locs=sorted(df['Location'].unique())
for k,l in enumerate(locs,2):
    s[f'A{k}']=l
    s[f'B{k}']=f'=COUNTIF({col("Location")},A{k})'
    for L,name in zip('CDEF',['Calories','Protein (g)','Carbohydrate (g)','Fat (g)']):
        s[f'{L}{k}']=f'=IFERROR(AVERAGEIFS({col(name)},{col("Location")},A{k},{col("Data Note")},""),0)'
t=len(locs)+2
s[f'A{t}']='All locations'; s[f'B{t}']=f'=SUM(B2:B{t-1})'
for L,name in zip('CDEF',['Calories','Protein (g)','Carbohydrate (g)','Fat (g)']):
    s[f'{L}{t}']=f'=AVERAGEIFS({col(name)},{col("Data Note")},"")'
for row in s.iter_rows(min_row=2,max_row=t):
    for c in row:
        c.font=F(bold=(c.row==t)); 
        if c.column>2: c.number_format='#,##0' if c.column==3 else '0.0'
s[f'A{t+2}']='Averages leave out rows with a Data Note (missing or likely wrong numbers).'; s[f'A{t+2}'].font=F(italic=True,color='6B6573')
s.column_dimensions['A'].width=26
for L in 'BCDEF': s.column_dimensions[L].width=16

a=wb.create_sheet('About')
lines=[('Maroon menu data',True),
('Source: McMaster Hospitality Services menu nutrition site, https://macnutrition.mcmaster.ca/Nutrition/ServiceMenuReport/Today',False),
('Updated automatically: a GitHub Action (.github/workflows/update-menu.yml) pulls the site every night and publishes changes. McMaster Hospitality confirmed (Sept 2026) that Maroon may use this data.',False),
('Nutrition values are exactly as posted by McMaster, per the listed serving size. For allergens, the site says to ask a Hospitality Manager.',False),
('',False),
('Cleanup done',True),
('Serving sizes with spelling or capitalization differences were merged (for example "portion", "PoRTION" and "Piortion" are now "Portion").',False),
('Thousands separators were removed so every nutrition column is a number.',False),
('Exact duplicate rows from the site were removed. Rows that share a name but have different numbers were kept.',False),
('',False),
('Data Note column (highlighted rows)',True),
('No nutrition info listed: the site shows the item with blank values. These are left out of the app.',False),
('Likely error in source data: numbers that cannot be right for one serving (for example a pulled pork taco listed at over 30,000 calories). These are left out of the app until McMaster fixes them.',False),
('',False),
('Other Restaurants tab',True),
('Tim Hortons (La Piazza): Tim Hortons Canada Nutrition Information, August 2025. Drinks, drink add-ons and baked goods only, since this location does not sell food.',False),
('Second Cup (PGCLL and Bistro 2 Go): Second Cup Beverage Menu Nutritional Values. Two blocks have no label in the PDF and are listed as Espresso and Americano by their position on the menu.',False),
('Chopped Leaf (PGCLL): Chopped Leaf Nutritional Chart. Proteins and dressings are add-ons in the app.',False),
('Paramount Lebanese Kitchen (Centro): calories from the menu boards (Sept 28, 2026). Paramount publishes no protein, carb or fat numbers, so those are ESTIMATES, as are the salad calories. Build-your-own dishes = posted base dish + posted protein.',False),
('Booster Juice (La Piazza and DBAC): Booster Juice Nutrition Guide, Version 24.1 (in-store sheet, photographed Sept 29, 2026). Wraps are on tomato tortillas; whole wheat adds 20 Cal.',False),
('Pizza Pizza (MUSC): pizzapizza.ca nutrition tables (facts as of Nov 1, 2023), pulled Oct 1, 2026. Classic pizzas only: walk-in slice and one slice of a Small/Medium/Large/X-Large. Sat fat printed without its decimal on the site was corrected.',False)]
for i,(t_,b) in enumerate(lines,1):
    a[f'A{i}']=t_; a[f'A{i}'].font=F(bold=b,size=13 if i==1 else 10); a[f'A{i}'].alignment=Alignment(wrap_text=True,vertical='top')
a.column_dimensions['A'].width=110
wb.move_sheet('About',offset=-3)
wb.save('mcmaster-menu-nutrition.xlsx')
print('ok',n)
